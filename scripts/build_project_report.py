"""Build the consolidated DOCX from reviewed Markdown; requires python-docx.

Run from any directory: python scripts/build_project_report.py
"""
from pathlib import Path
import re
from datetime import datetime
from zoneinfo import ZoneInfo
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SECTIONS = [
    ('1. Review, architecture and assessment', 'docs/REVIEW.md'),
    ('2. Docker installation and operation', 'docs/DOCKER.md'),
    ('3. Current execution evidence', 'docs/VALIDATION.md'),
    ('4. Detailed project guide', 'docs/PROJECT_GUIDE.md'),
    ('5. Dataset processing and implementation logic', 'docs/DATA_PROCESSING_AND_PROJECT_LOGIC.md'),
    ('6. Archived research measurements', 'reports/research_results.md'),
    ('7. Dataset provenance and source revisions', 'docs/DATASET_SOURCES.md'),
]


def inline(paragraph, text):
    # Preserve link destinations in a portable, readable report.
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    for token in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        run = paragraph.add_run(token[2:-2] if token.startswith('**') else token[1:-1] if token.startswith('`') else token)
        if token.startswith('**'):
            run.bold = True
        if token.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(9)


def render(doc, text):
    lines = text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if not line.strip():
            i += 1
            continue
        if line.startswith('```'):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i]); i += 1
            # Individual paragraphs permit long code blocks to break across pages.
            for code_line in code:
                p = doc.add_paragraph(style='Code Block')
                p.add_run(code_line)
            i += 1
            continue
        if line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [v.strip() for v in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', c.replace(' ', '')) for c in cells):
                    rows.append(cells)
                i += 1
            table = doc.add_table(rows=0, cols=max(map(len, rows)))
            table.style = 'Light Shading Accent 1'
            for n, row in enumerate(rows):
                cells = table.add_row().cells
                for c, value in zip(cells, row):
                    inline(c.paragraphs[0], value)
                    for run in c.paragraphs[0].runs:
                        run.font.size = Pt(8.5)
                        if n == 0: run.bold = True
                if n == 0:
                    repeat = OxmlElement('w:tblHeader')
                    table.rows[0]._tr.get_or_add_trPr().append(repeat)
            doc.add_paragraph()
            continue
        if match := re.match(r'^(#{1,6})\s+(.*)', line):
            doc.add_heading(match[2], level=min(len(match[1]) + 1, 4))
        elif re.match(r'^\s*[-*]\s+', line):
            inline(doc.add_paragraph(style='List Bullet'), re.sub(r'^\s*[-*]\s+', '', line))
        elif re.match(r'^\d+\.\s+', line):
            inline(doc.add_paragraph(style='List Number'), re.sub(r'^\d+\.\s+', '', line))
        else:
            inline(doc.add_paragraph(), line)
        i += 1


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.top_margin = sec.bottom_margin = Inches(.75)
    sec.left_margin = sec.right_margin = Inches(.75)
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for level in range(1, 5):
        style = doc.styles[f'Heading {level}']
        style.font.color.rgb = RGBColor.from_string('17365D')
    code = doc.styles.add_style('Code Block', 1)
    code.font.name = 'Consolas'; code.font.size = Pt(8)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.left_indent = Inches(.12)
    code.paragraph_format.line_spacing = 1
    sec.header.paragraphs[0].text = 'PROMPTGUARD  |  PROJECT REVIEW & TECHNICAL REPORT'
    sec.header.paragraphs[0].runs[0].font.size = Pt(8)
    footer = sec.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('PromptGuard Research • Page ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.add_paragraph('TECHNICAL PROJECT REPORT', 'Subtitle')
    doc.add_heading('Prompt Injection Detection Research', 0)
    doc.add_paragraph('Architecture • Data provenance • Machine learning methodology\nCode review • Docker operation • Validation • Research limitations', 'Subtitle')
    doc.add_paragraph('Repository: Sumukha-Rao/Minor-Project\nReviewed revision: cbf6720783ca02240ed14c3c7c9ae54aeddc5d86\nIncludes uncommitted containerization and documentation changes.')
    doc.add_paragraph('Generated: ' + datetime.now(ZoneInfo('Asia/Kolkata')).strftime('%d %B %Y, %H:%M %Z'))
    doc.add_heading('How to read this report', 1)
    doc.add_paragraph('Chapters 1–3 describe this delivery and its fresh validation. Chapters 4–7 consolidate the repository documentation and historical evidence. Historical local paths and hardware are not requirements for the Docker workflow. Prior research measurements and tiny-model tests are never presented as new real Transformer experiments.')
    doc.add_paragraph('The application is a Python command-line research pipeline. Docker provides a CPU execution environment and persistent output directories; it does not turn the detector into a web service or establish production security effectiveness.')
    doc.add_page_break()
    doc.add_heading('Contents', 1)
    for title, _ in SECTIONS: doc.add_paragraph(title)
    doc.add_paragraph('Use Word/LibreOffice’s navigation pane for the nested headings. Page numbers are updated by the document viewer.')
    doc.add_heading('Evidence legend', 2)
    for text in ['Implemented: supported by inspected source code.', 'Executed: supported by current container logs and artifacts in Chapter 3.', 'Archived: recorded by the repository’s previous experiment, with explicit source attribution.', 'Pending: not executed or dependent on unavailable pretrained assets.']:
        doc.add_paragraph(text, 'List Bullet')
    for title, filename in SECTIONS:
        doc.add_page_break()
        doc.add_heading(title, 1)
        doc.add_paragraph('Source: ' + filename, 'Caption')
        if title.startswith(('4.', '5.', '6.')):
            doc.add_paragraph('Historical repository documentation follows. Availability statements and machine-specific measurements refer to the original workspace. Current commands and execution outcomes are in Chapters 2 and 3.')
        render(doc, (ROOT / filename).read_text())
    doc.core_properties.title = 'PromptGuard — Comprehensive Project Review and Docker Guide'
    doc.core_properties.subject = 'Architecture, datasets, ML methods, containerization, validation and limitations'
    doc.core_properties.author = 'Project technical documentation'
    output = ROOT / 'docs/PROJECT_REPORT.docx'
    doc.save(output)
    # Structural round-trip check catches malformed/unreadable OOXML packages.
    check = Document(output)
    print(f'{output}: {len(check.paragraphs)} paragraphs, {len(check.tables)} tables, {output.stat().st_size:,} bytes')


if __name__ == '__main__':
    main()
